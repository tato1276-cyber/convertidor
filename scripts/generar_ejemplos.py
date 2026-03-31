from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


def generar_csv(path: Path) -> None:
    contenido = "\n".join(
        [
            "codigo,sku,descripcion,brand,modelo,rubro,precio_sin_iva,iva,stock,estado,moneda,link,nota",
            "A1,SKU-A1,Notebook 14,MarcaX,M14,Notebooks,1000,21,5,disponible,USD,http://example.com/a1,ok",
        ]
    )
    path.write_text(contenido, encoding="utf-8")


def generar_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Lista"
    ws.append(
        [
            "codigo",
            "sku",
            "descripcion",
            "brand",
            "modelo",
            "rubro",
            "precio_sin_iva",
            "iva",
            "stock",
            "estado",
            "moneda",
            "link",
            "nota",
        ]
    )
    ws.append(
        [
            "A4",
            "SKU-A4",
            "Monitor 24",
            "MarcaX",
            "M24",
            "Monitores",
            200,
            21,
            3,
            "disponible",
            "USD",
            "http://example.com/a4",
            "ok",
        ]
    )
    wb.save(path)


def generar_pdf(path: Path) -> None:
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont("Helvetica", 11)
    lineas = [
        "codigo | descripcion | precio_final",
        "A2 | Mouse USB | 242",
        "A5 | Webcam HD | 363",
    ]
    y = 750
    for linea in lineas:
        c.drawString(72, y, linea)
        y -= 18
    c.save()


def generar_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("codigo;descripcion;precio_final")
    doc.add_paragraph("A3;Teclado;363")
    doc.add_paragraph("A6;Parlantes;484")
    doc.save(path)


def main() -> None:
    base = Path(__file__).resolve().parents[1] / "ejemplos"
    base.mkdir(parents=True, exist_ok=True)

    generar_csv(base / "proveedor_alpha.csv")
    generar_xlsx(base / "proveedor_alpha.xlsx")
    generar_pdf(base / "proveedor_alpha.pdf")
    generar_docx(base / "proveedor_alpha.docx")

    print(f"Ejemplos generados en: {base}")


if __name__ == "__main__":
    main()
