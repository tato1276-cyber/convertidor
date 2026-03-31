from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from docx import Document


def _clean_text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _extract_tables(doc: Document) -> list[pd.DataFrame]:
    tables: list[pd.DataFrame] = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [_clean_text(cell.text) for cell in row.cells]
            if any(cells):
                rows.append(cells)

        if len(rows) < 2:
            continue

        header = rows[0]
        width = len(header)
        body = []
        for row in rows[1:]:
            padded = row + [""] * max(0, width - len(row))
            body.append(padded[:width])
        tables.append(pd.DataFrame(body, columns=header))
    return tables


def _extract_paragraph_rows(doc: Document) -> pd.DataFrame:
    parsed = []
    for para in doc.paragraphs:
        line = _clean_text(para.text)
        if not line:
            continue

        for pattern in [r"\s*\|\s*", r"\s*;\s*", r"\t+", r"\s{2,}"]:
            parts = re.split(pattern, line)
            parts = [p.strip() for p in parts if p.strip()]
            if len(parts) >= 3:
                parsed.append(parts)
                break

    if len(parsed) < 2:
        raise ValueError("DOCX sin tabla interpretable: no se detectaron columnas suficientes.")

    header = parsed[0]
    width = len(header)
    body = []
    for row in parsed[1:]:
        padded = row + [""] * max(0, width - len(row))
        body.append(padded[:width])

    return pd.DataFrame(body, columns=header)


def read_docx_file(file_path: Path) -> pd.DataFrame:
    """Extrae tablas simples de .docx. Si no hay tablas, intenta por párrafos delimitados."""
    doc = Document(file_path)
    tables = _extract_tables(doc)
    if tables:
        return pd.concat(tables, ignore_index=True)

    return _extract_paragraph_rows(doc)
